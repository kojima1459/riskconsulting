#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
bin_roundtrip.py - 配布 vbaProject.bin の読み戻し検問(裁定書27 W9-A 検証① /
                   新ゲート `bin-roundtrip`)

================================================================================
何を守るゲートか:
    配布方式B(裁定書27 W9-A)は「モジュールが最初から入った正規の
    vbaProject.bin をビルドが書き出す」方式である。書き出す側(build/ovba_write.py)
    が間違っても、書いた側の理屈で読み返しては何も検査したことにならない。
    本ツールは**配布物(dist/*.xlsm)の側から**bin を開き、

      [1] 各モジュールのソースが src/ の .bas/.cls(ヘッダ除去・CRLF・CP932)と
          **バイト一致**すること
      [2] モジュール数と集合が build/modules.json(その配布に載る分)と一致すること
      [3] 隠しシート vba_src が存在しないこと
      [4] bin に配布禁止の文字列("VBProject" / "AddFromString" /
          "ExecuteExcel4Macro" / "WScript.Shell" / "new:{")が現れないこと
          (裁定書27 W9-B 6。**圧縮を解いた本文で**検査する)
      [5] 全モジュールの MODULEOFFSET が 0(p-code キャッシュを持たない)こと
      [6] モジュールの「形」が Mac 実Excel 製ブックと一致すること(W9.3)
          - class : 属性8行(VB_Base = 0{FCFB3D2A-A0FA-1068-A738-08002B3371B5})
                    ＋ dir の MODULE レコードに MODULEPRIVATE(0x0028・Size=0)
          - document: 0x0028 は**無い**・VB_Customizable = True
          - std   : 属性は `Attribute VB_Name` の1行だけ

    を確かめる。読めない・数えられない・比較できないは**すべて失格**にする
    (「対象が見つからないので検査せず緑」を作らない)。

なぜ olevba を使うのか:
    [1] の解凍は oletools.olevba(第三者実装)で行う。ビルドが使った
    build/ovba.py の解凍器で読み返すと「自分の圧縮器のバグを自分の解凍器が
    帳消しにする」ため、往復検査としての意味が消える。
    olevba が入っていない場合は**緑にせず** exit 2(環境不備)で止める。

整形規則の値源:
    .bas → モジュールソースの整形(Attribute行/.clsヘッダの除去・改行正規化)は
    build/build_rpn.py の `_vba_src_text` **1実装だけ**を呼ぶ。ここで書き写すと
    ビルドと検問が別々に緩められるため、必ず import して使う。

突合の根拠(W9.3・2026-09-03):
    [6] の期待値は **Mac の実Excel が保存したブック**(Class1 を1本足しただけの
    サンプル。scratchpad/bisect/mac_class_sample.xlsm)の dir と各モジュール
    ストリームを解析して得た実測値である。仕様書([MS-OVBA])だけでは
    「Excel が実際に何を書くか」が決まらず、我々の bin は
      (1) クラスの MODULE レコードに MODULEPRIVATE(0x0028)が無い
      (2) クラスの属性行が5行しかない(VB_Base / VB_TemplateDerived /
          VB_Customizable を欠く)
    という2点で Excel と食い違っており、**クラスを1本含めるだけで Mac の実Excel が
    読み込み時に「実行時エラー 5」の生ダイアログを出していた**(17章 Z-24)。
    両方を直した版が実機で正常に開くことを確認済み。この検問はその退行を止める。
    なお同サンプルは**全モジュールの改行が LF 単独**だった(Windows 製は CRLF)。
    我々は CRLF のままにしている(CRLF の切り分けブックが実機で通っており、
    改行は Err 5 の要因ではない)。事実として記録に残す。

使い方:
    python3 tools/bin_roundtrip.py                 # dist/ の dev と prod 両方
    python3 tools/bin_roundtrip.py --book dist/リスク提案ナビ.xlsm
    exit code: 0 = 全PASS / 1 = 失格 / 2 = 環境不備(oletools不在・ブック不在)
================================================================================
"""

from __future__ import annotations

import argparse
import sys
import zipfile
from pathlib import Path

TOOLS_DIR = Path(__file__).resolve().parent
REPO_ROOT = TOOLS_DIR.parent
sys.path.insert(0, str(REPO_ROOT / "build"))

import build_rpn      # noqa: E402  (_vba_src_text / 禁止文字列表 / 台帳読み)
import ovba_write     # noqa: E402  (dir の申告を読む)

try:
    from oletools.olevba import VBA_Parser
except ImportError:   # 第三者実装が無いと往復検査にならないので緑にしない
    VBA_Parser = None

DEFAULT_BOOKS = (
    REPO_ROOT / "dist" / "リスク提案ナビ_dev.xlsm",
    REPO_ROOT / "dist" / "リスク提案ナビ.xlsm",
)
DEV_SUFFIX = "_dev.xlsm"
# 第2段の一段化(17章 Z-42)の産物 dist/final/ にだけ載る designer モジュール。
# --final を付けたときだけ「居てよい」ものとして数える(既定の第1段検査は不変)。
FINAL_FORM_NAME = "frmNaviHtml"
FINAL_FORM_FRM = REPO_ROOT / "src" / "ui" / "navi" / "frmNaviHtml.frm"
# designer(UserForm)の VB_Base は「そのフォーム固有の2つのGUID」であり、
# クラスモジュールの固定値(CLASS_VB_BASE)とは別物。値源は build/ovba_write.py。
FINAL_FORM_VB_BASE = 'Attribute VB_Base = "%s"' % ovba_write.FORM_VB_BASE


def expected_sources(is_dev: bool, final: bool = False) -> dict:
    """台帳から {モジュール名: 期待するモジュール本文(bytes)} を作る。"""
    modules = build_rpn.load_manifest(str(REPO_ROOT / "build" / "modules.json"))
    # モード別ソース選択(裁定書30 裁定1(b))。dev ブックの modGatewayLink /
    # modTestsPureHook は dev_src のソースと突き合わせる。
    modules = build_rpn.select_variant_paths(modules, is_dev)
    present, _missing = [], []
    for m in modules:
        if (REPO_ROOT / m["path"]).exists():
            present.append(m)
    shipped = build_rpn._shipped_modules(present, is_dev)
    out = {}
    for m in shipped:
        body = build_rpn._vba_src_text(str(REPO_ROOT), m)
        want = body.replace("\n", "\r\n").encode("cp932")
        if not want.endswith(b"\r\n"):
            want += b"\r\n"
        out[m["name"]] = want
    if final:
        # UserForm は modules.json で type=form(第1段に載せない)。final の bin
        # にだけ居るので、期待本文は .frm の Begin…End より後ろのコード部から作る。
        with open(FINAL_FORM_FRM, encoding="utf-8") as fh:
            _d, _a, code = ovba_write.split_frm(fh.read())
        want = "\n".join(code).replace("\n", "\r\n").encode("cp932")
        if not want.endswith(b"\r\n"):
            want += b"\r\n"
        out[FINAL_FORM_NAME] = want
    return out


def check_book(book: Path, final: bool = False) -> list[str]:
    errors: list[str] = []
    is_dev = book.name.endswith(DEV_SUFFIX)
    print("=" * 78)
    print(f"ブック: {book}  ({'dev' if is_dev else 'prod'})")
    print("=" * 78)

    with zipfile.ZipFile(book) as z:
        names = z.namelist()
        if "xl/vbaProject.bin" not in names:
            return [f"{book.name}: xl/vbaProject.bin がありません"]
        vba_bin = z.read("xl/vbaProject.bin")

    # --- [3] vba_src シートの不在 -------------------------------------------
    import openpyxl
    wb = openpyxl.load_workbook(book, read_only=True, keep_links=False)
    sheetnames = list(wb.sheetnames)
    wb.close()
    if "vba_src" in sheetnames:
        errors.append(f"{book.name}: 隠しシート vba_src が残っています"
                      "(配布方式Bでは存在してはいけません)")
    print(f"[3] vba_src シート: {'あり(失格)' if 'vba_src' in sheetnames else 'なし'}"
          f" / シート{len(sheetnames)}枚")

    # --- [1] olevba で解凍して .bas とバイト比較 ------------------------------
    got: dict[str, bytes] = {}
    parser = VBA_Parser(str(book))
    try:
        for (_fn, stream, _vba_fn, code) in parser.extract_macros():
            name = stream.split("/")[-1]
            got[name] = code.encode("cp932", errors="replace") \
                if isinstance(code, str) else code
    finally:
        parser.close()
    print(f"[1] olevba が解凍したモジュール: {len(got)}本")

    want = expected_sources(is_dev, final)
    doc_names = {n for n, info in ovba_write.read_modules(vba_bin).items()
                 if info["type"] == "document"}
    # 配布方式Bで焼く document module は ThisWorkbook のみ(build_rpn.py
    # build_baked_vba_project)。Sheet1 等の他の document module は焼かない
    # (openpyxl 製の成果物ワークシートに codeName="Sheet1" が無く、焼くと
    # 名前だけの孤児モジュールになるため)。ここで document 集合が
    # {"ThisWorkbook"} ちょうどであることを確かめる(集合一致条件は緩めない)。
    if doc_names != {"ThisWorkbook"}:
        errors.append(
            f"{book.name}: document module 集合が {{'ThisWorkbook'}} と不一致"
            f"(実際: {sorted(doc_names)})。配布方式Bは ThisWorkbook 以外の "
            "document module を焼かない設計です。")
    std_got = {n: v for n, v in got.items() if n not in doc_names}

    for name, want_src in sorted(want.items()):
        if name not in std_got:
            errors.append(f"{book.name}: '{name}' が配布binにありません")
            continue
        actual = ovba_write.strip_attribute_lines(std_got[name])
        if actual != want_src:
            errors.append(
                f"{book.name}: '{name}' の本文が src/ と不一致"
                f"(期待{len(want_src)}バイト / 実際{len(actual)}バイト)")
    extra = sorted(set(std_got) - set(want))
    if extra:
        errors.append(f"{book.name}: 台帳に無いモジュールが載っています: {extra}")

    # --- [2] モジュール数と集合 ----------------------------------------------
    print(f"[2] モジュール集合: 台帳{len(want)}本 / bin(document除く){len(std_got)}本"
          f" / document module {sorted(doc_names)}")
    if len(std_got) != len(want):
        errors.append(
            f"{book.name}: モジュール数が台帳と不一致(台帳{len(want)} / bin{len(std_got)})")

    # --- [4] 禁止文字列 -------------------------------------------------------
    # prod ブックは direct経路の痕跡(ServerXMLHTTP 等)も見る(裁定書30 裁定1(f))。
    hits = build_rpn.forbidden_strings_in_bin(vba_bin, prod=not is_dev)
    print(f"[4] 配布禁止文字列: {hits if hits else 'なし'}")
    if hits:
        errors.append(
            f"{book.name}: vbaProject.bin に配布禁止の文字列があります"
            f"(裁定書27 W9-B 6): {', '.join(hits)}")

    # --- [4b] 起動スタブの形(W9.2) ------------------------------------------
    # ThisWorkbook の Workbook_Open は modBoot.Boot を直接呼ぶ。ブック名で
    # 修飾した Application.Run(非ASCIIブック名をExcel側に解決させる形)は
    # Mac実機で Err 5 の生ダイアログを出した(裁定書27 W9.2)。スタブ本体は
    # bin-roundtrip の「ソース一致」では検出できない(スタブ自体が値源)ため、
    # ここで意味として禁じる。
    mods_info = ovba_write.read_modules(vba_bin)
    tw = mods_info.get("ThisWorkbook", {}).get("source", b"")
    tw_text = tw.decode("cp932", errors="replace") if isinstance(tw, (bytes, bytearray)) else str(tw)
    bad_stub = [k for k in ("ThisWorkbook.Name", "Application.Run", "VBProject") if k in tw_text]
    print(f"[4b] ThisWorkbook スタブの禁止形: {bad_stub if bad_stub else 'なし'}")
    if bad_stub:
        errors.append(
            f"{book.name}: ThisWorkbook スタブに禁止の形があります(W9.2): {', '.join(bad_stub)}")

    # ブックイベントの受け口は **ThisWorkbook 文書モジュール**である(W9.3)。
    # 旧実装は WithEvents を持つクラス(clsAppEvents)で受けていたが、配布物から
    # クラスモジュールを外した(可動部品を減らす。17章 Z-24)。その結果、
    # 「全画面を当て直す/元へ戻す/閉じるときトーストの予約を取り消す」経路は
    # このスタブにしか存在しない。**焼き忘れても他のどの検問にも引っかからない**
    # (スタブ自体が値源であり src/ に対応物が無い)ので、ここで名指しで見る。
    REQUIRED_STUB_SUBS = ("Workbook_Open", "Workbook_Activate",
                          "Workbook_Deactivate", "Workbook_BeforeClose")
    missing_stub = [k for k in REQUIRED_STUB_SUBS if k not in tw_text]
    print(f"[4b] ThisWorkbook スタブの必須イベント: "
          f"{'すべてあり' if not missing_stub else '欠落=' + str(missing_stub)}")
    if missing_stub:
        errors.append(
            f"{book.name}: ThisWorkbook スタブに必須のブックイベントがありません"
            f"(W9.3): {', '.join(missing_stub)}")

    # 配布物にクラスモジュールが無いこと(W9.3)。document は ThisWorkbook 1本。
    cls_names = sorted(n for n, i in mods_info.items() if i["type"] == "class")
    print(f"[4b] クラスモジュール: {cls_names if cls_names else '0本'}")
    if final:
        # designer(UserForm)は dir 上はクラスと同型(0x0022 + MODULEPRIVATE)。
        # --final のときだけ frmNaviHtml を除いて数える。他のクラスは依然禁止。
        cls_names = [n for n in cls_names if n != FINAL_FORM_NAME]
    if cls_names:
        errors.append(
            f"{book.name}: 配布物にクラスモジュールが載っています(W9.3 の配布方針"
            f"ではブックイベントは ThisWorkbook が受け、クラスは持ちません): {cls_names}")

    # --- [5] MODULEOFFSET=0 ---------------------------------------------------
    import struct
    import ovba
    dir_dec = ovba.ovba_decompress(ovba.CFBReader(vba_bin).read("dir"))
    offsets = [struct.unpack("<I", body)[0]
               for _o, rid, _s, body in ovba_write.iter_dir_records(dir_dec)
               if rid == ovba_write.REC_MODULEOFFSET]
    bad = [o for o in offsets if o != 0]
    print(f"[5] MODULEOFFSET: {len(offsets)}件すべて{'0' if not bad else '0ではない'}"
          "(p-codeキャッシュ無し)")
    if not offsets:
        errors.append(f"{book.name}: dir に MODULEOFFSET が1件もありません"
                      "(検査が成立していません)")
    if bad:
        errors.append(f"{book.name}: MODULEOFFSET≠0 のモジュールが{len(bad)}件"
                      "(p-codeキャッシュが混入しています)")

    # --- [6] モジュールの「形」が Mac 実Excel 製と一致するか(W9.3) -----------
    errors.extend(check_module_shapes(
        book.name, vba_bin, dir_dec, mods_info,
        designer_names=(FINAL_FORM_NAME,) if final else ()))

    # --- [7] 一段化(Z-42)の産物だけの条件: 参照設定と designer ストレージ ---
    if final:
        refs = ovba_write.read_reference_names(vba_bin)
        print(f"[7] 参照設定: {refs}")
        for need in ("SHDocVw", "MSForms"):
            if need not in refs:
                errors.append(f"{book.name}: 参照設定 {need} がありません")
        for r in refs:
            if r not in ("VBA", "Excel", "stdole", "Office", "SHDocVw", "MSForms"):
                errors.append(f"{book.name}: 許可していない参照設定があります: {r}")
        try:
            st = ovba_write.read_designer_storage(vba_bin, FINAL_FORM_NAME)
        except ovba_write.OvbaWriteError as e:
            st = {}
            errors.append(f"{book.name}: {e}")
        print(f"[7] designer ストレージ {FINAL_FORM_NAME}/: "
              f"{sorted(repr(k) for k in st)}")
        for nm in ovba_write.FORM_STREAM_NAMES:
            if nm not in st:
                errors.append(
                    f"{book.name}: designer ストリーム {FINAL_FORM_NAME}/{nm!r} が"
                    "ありません(フォームが開けません)")
        if st.get("\x03VBFrame") and not st["\x03VBFrame"].startswith(b"VERSION "):
            errors.append(f"{book.name}: \\x03VBFrame がデザイナ定義になっていません")
    return errors


# 実測値(Mac 実Excel 製サンプルの Class1)。ここを緩めると Z-24 が再発する。
CLASS_VB_BASE = 'Attribute VB_Base = "0{FCFB3D2A-A0FA-1068-A738-08002B3371B5}"'
REC_MODULEPRIVATE = 0x0028


def module_private_flags(dir_dec: bytes) -> dict[str, bool]:
    """dir を1回舐めて {モジュール名: MODULEPRIVATE(0x0028)を持つか} を返す。

    MODULENAME(0x0019)で名前が確定し、MODULE_TERMINATOR(0x002B)で1本が閉じる。
    その間に 0x0028 が現れたかどうかを記録する。
    """
    out: dict[str, bool] = {}
    name = None
    private = False
    for _off, rid, _sz, body in ovba_write.iter_dir_records(dir_dec):
        if rid == ovba_write.REC_MODULENAME:
            name = body.decode("cp932", errors="replace")
            private = False
        elif rid == REC_MODULEPRIVATE:
            private = True
        elif rid == ovba_write.REC_MODULE_TERMINATOR and name is not None:
            out[name] = private
            name = None
    return out


def check_module_shapes(book_name: str, vba_bin: bytes, dir_dec: bytes,
                        mods_info: dict, verbose: bool = True,
                        designer_names: tuple = ()) -> list[str]:
    """[6] 属性行と dir の MODULEPRIVATE が Mac 実Excel 製と同じ形か。"""
    errors: list[str] = []
    privates = module_private_flags(dir_dec)
    counts = {"std": 0, "class": 0, "document": 0}
    for name, info in sorted(mods_info.items()):
        src = info["source"]
        text = src.decode("cp932", errors="replace") if isinstance(src, (bytes, bytearray)) else str(src)
        # 先頭の連続する属性行だけを見る(本文中のメンバー属性
        # `Attribute App.VB_VarHelpID = -1` を巻き込まない)。
        header = []
        for ln in text.split("\r\n"):
            if not ln.startswith("Attribute "):
                break
            header.append(ln)
        # ovba_write.read_modules は 0x0021 を "procedural" と呼ぶ。
        kind = {"procedural": "std"}.get(info["type"], info["type"])
        counts[kind] = counts.get(kind, 0) + 1
        has_private = privates.get(name)
        if has_private is None:
            errors.append(f"{book_name}: '{name}' が dir の MODULE レコードに見つかりません")
            continue
        if kind == "std":
            if len(header) != 1 or not header[0].startswith('Attribute VB_Name'):
                errors.append(
                    f"{book_name}: 標準モジュール '{name}' の属性行が "
                    f"`Attribute VB_Name` の1行だけではありません(実際{len(header)}行)")
            if has_private:
                errors.append(f"{book_name}: 標準モジュール '{name}' に "
                              "MODULEPRIVATE(0x0028)があります(Excel は付けません)")
        elif kind == "class" and name in designer_names:
            # designer(UserForm)。属性8行は同じだが VB_Base はフォーム固有。
            if len(header) != 8:
                errors.append(
                    f"{book_name}: designer '{name}' の属性行が8行では"
                    f"ありません(実際{len(header)}行)")
            if FINAL_FORM_VB_BASE not in header:
                errors.append(
                    f"{book_name}: designer '{name}' に "
                    f"`{FINAL_FORM_VB_BASE}` がありません")
            if not has_private:
                errors.append(
                    f"{book_name}: designer '{name}' の dir に "
                    "MODULEPRIVATE(0x0028)がありません")
        elif kind == "class":
            if len(header) != 8:
                errors.append(
                    f"{book_name}: クラスモジュール '{name}' の属性行が8行では"
                    f"ありません(実際{len(header)}行。Mac実Excel製と不一致)")
            if CLASS_VB_BASE not in header:
                errors.append(
                    f"{book_name}: クラスモジュール '{name}' に "
                    f"`{CLASS_VB_BASE}` がありません")
            if not has_private:
                errors.append(
                    f"{book_name}: クラスモジュール '{name}' の dir に "
                    "MODULEPRIVATE(0x0028)がありません(17章 Z-24 の再発)")
        elif kind == "document":
            if has_private:
                errors.append(f"{book_name}: document module '{name}' に "
                              "MODULEPRIVATE(0x0028)があります(Excel は付けません)")
            if "Attribute VB_Customizable = True" not in header:
                errors.append(
                    f"{book_name}: document module '{name}' に "
                    "`Attribute VB_Customizable = True` がありません")
    if verbose:
        print(f"[6] モジュールの形: std{counts.get('std', 0)}本 / "
              f"class{counts.get('class', 0)}本 / document{counts.get('document', 0)}本"
              f"(属性行と MODULEPRIVATE(0x0028)を Mac実Excel製と突合)")
    return errors


SELFTEST_CLASS_BODY = (
    "Option Explicit\r\n"
    "Public WithEvents App As Application\r\n"
    "Attribute App.VB_VarHelpID = -1\r\n"
)


def selftest_class_shape() -> list[str]:
    """[6] のクラス規則が**空振りしていない**ことの自己テスト(W9.3)。

    配布物(リスク提案ナビ)はクラスモジュールを持たない(ブックイベントは
    ThisWorkbook が受ける)ため、上の [6] のクラス分岐は配布物を検査するだけでは
    **一度も走らない**。走らない検査は数か月で腐る(build/ovba_write.py の
    クラス対応は他プロダクトも使う)。そこで、その場でクラス入りの bin を1本
    組み立てて、
      (正例) 属性8行 + MODULEPRIVATE(0x0028)を書いた bin は 0 件で通ること
      (負例) dir から 0x0028 を抜いた bin は必ず落ちること
    を毎回確かめる。負例が通ってしまう=検査が骨抜きになった、である。
    """
    import ovba
    tmpl_path = REPO_ROOT / "build" / "template_skeleton.xlsm"
    if not tmpl_path.exists():
        return ["自己テスト: build/template_skeleton.xlsm がありません"]
    with zipfile.ZipFile(tmpl_path) as z:
        tmpl_bin = z.read("xl/vbaProject.bin")

    mods = [
        ovba_write.VbaModule(
            "ThisWorkbook",
            ovba_write.module_stream_source("ThisWorkbook", "Option Explicit\r\n",
                                            "document"),
            "document"),
        ovba_write.VbaModule(
            "modSelfTest",
            ovba_write.module_stream_source("modSelfTest", "Option Explicit\r\n", "std"),
            "std"),
        ovba_write.VbaModule(
            "clsSelfTest",
            ovba_write.module_stream_source("clsSelfTest", SELFTEST_CLASS_BODY, "class"),
            "class"),
    ]
    vba_bin = ovba_write.build_vba_project(tmpl_bin, mods)
    mods_info = ovba_write.read_modules(vba_bin)
    dir_dec = ovba.ovba_decompress(ovba.CFBReader(vba_bin).read("dir"))

    out: list[str] = []
    pos = check_module_shapes("(自己テスト・正例)", vba_bin, dir_dec, mods_info,
                              verbose=False)
    if pos:
        out.append("自己テスト(正例): 属性8行+0x0028 を書いた bin が [6] で落ちました: "
                   + " / ".join(pos))

    # 負例: dir から MODULEPRIVATE(0x0028)のレコードを1本残らず抜く。
    stripped = bytearray()
    for off, rid, _sz, body in ovba_write.iter_dir_records(dir_dec):
        if rid == REC_MODULEPRIVATE:
            continue
        # 実長は body の長さで取る(PROJECTVERSION は Size フィールドの申告 4 に
        # 対して実データが 6 バイトある既知の例外。申告値で切ると dir が壊れる)。
        stripped += dir_dec[off:off + 6 + len(body)]
    neg = check_module_shapes("(自己テスト・負例)", vba_bin, bytes(stripped), mods_info,
                              verbose=False)
    if not any("MODULEPRIVATE" in m for m in neg):
        out.append("自己テスト(負例): dir から 0x0028 を抜いた bin を [6] が"
                   "見逃しました(クラス規則が骨抜きになっています)")
    return out


def main() -> int:
    ap = argparse.ArgumentParser(
        description="配布 vbaProject.bin の読み戻し検問(裁定書27 W9-A)")
    ap.add_argument("--book", help="検査するブック(既定: dist/ の dev と prod 両方)")
    ap.add_argument("--final", action="store_true",
                    help="第2段の一段化(17章 Z-42)の産物を検査する。"
                         "UserForm frmNaviHtml と参照設定(SHDocVw/MSForms)と "
                         "designer ストレージ4本が**在ること**を足して見る。"
                         "--book を省くと dist/final/リスク提案ナビ.xlsm")
    args = ap.parse_args()

    if VBA_Parser is None:
        print("ERROR: oletools が import できません(pip install oletools)。"
              "第三者実装で解凍できないと往復検査になりません。", file=sys.stderr)
        return 2

    if args.book:
        p = Path(args.book)
        books = [p if p.is_absolute() else REPO_ROOT / p]
    elif args.final:
        books = [REPO_ROOT / "dist" / "final" / "リスク提案ナビ.xlsm"]
    else:
        books = [p for p in DEFAULT_BOOKS if p.exists()]
    if not books:
        print("ERROR: dist/ にビルド済みブックがありません。"
              "先に python3 build/build_rpn.py --dev を実行してください。",
              file=sys.stderr)
        return 2

    print("=== bin_roundtrip.py (配布binを解凍して src/ とバイト比較) ===")
    errors: list[str] = []
    st = selftest_class_shape()
    print(f"[6自己] クラス規則の自己テスト(正例/負例): {'OK' if not st else 'NG'}")
    errors.extend(st)
    for b in books:
        if not b.exists():
            print(f"ERROR: ブックが見つかりません: {b}", file=sys.stderr)
            return 2
        errors.extend(check_book(b, final=args.final))
        print()

    print("-" * 78)
    if errors:
        print(f"結果: NG {len(errors)}件")
        for e in errors:
            print(f"  - {e}")
        return 1
    print(f"結果: OK 検査したブック{len(books)}冊 / {'7' if args.final else '6'}条件"
          "(本文バイト一致・モジュール数と集合・vba_src不在・禁止文字列不在・"
          "MODULEOFFSET=0・モジュールの形がMac実Excel製と一致"
          + ("・参照設定と designer ストレージ" if args.final else "") + ")")
    return 0


if __name__ == "__main__":
    sys.exit(main())
