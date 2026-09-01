#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
sheet_check.py - 13章(データ設計)と、シート台帳/ビルド済みブックの照合(17章 T-02/T-03)

================================================================================
**13章が正・build/sheets_main.json が従・成果物ブックがその実体**という三者関係を、
13章のMarkdown表を直接パースして機械照合する。ずれたら台帳(JSON)を直す。
13章そのものの誤りを疑ったときは直さずに司令塔へ報告する(17章§1の規約)。

照合するもの(19章§5「13章のシート・列と11章の画面項目が一致」の照合手順に従う):

  1. シート集合   13章§2の節見出しとガードシート規定から本体18枚を導き、
                  JSON・ブックの双方と突合する(順序はJSONが正・ブックはJSON順)。
  2. シート型     13章§2.9の型表(帳票型=HOME/案件入力・テーブル型=6枚)と
                  JSONの role が一致するか。
  3. 列物理名     テーブル型シートの列名と物理順。1シート1テーブルのシートは
                  ブックの1行目、ブロック縦積みのシートは**ブロック名の名前付き
                  レンジ(アンカー)から特定したヘッダ行**を読む(13章§2.9)。
  4. 名前付きレンジ 帳票型2シート(hm_ / ci_)と、ヒアリングシート・壁打ちの
                  見出し用(hs_ / sp_ の計8本)を名前付きレンジ側で突合する。
  5. configキー   13章§2.3の name 列と、その順序。機械比較できる行は既定値も。

ビルド機構シート(sheets_main.json で build_infrastructure=true のもの。vba_src)は
13章に存在しないため、**黙って無視せず**「仕様外のビルド機構」として明示的に
除外し、その事実を出力する。

使い方:
    python3 tools/sheet_check.py
    python3 tools/sheet_check.py --book dist/リスク提案ナビ.xlsm
    python3 tools/sheet_check.py --no-book      # 台帳(JSON)と13章だけを突合する
    exit code: 0 = 全一致 / 1 = 不一致あり
================================================================================
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

TOOLS_DIR = Path(__file__).resolve().parent
REPO_ROOT = TOOLS_DIR.parent
SPEC_CH13 = REPO_ROOT / "docs" / "spec" / "13_データ設計.md"
SPEC_CH19 = REPO_ROOT / "docs" / "spec" / "19_用語集とレジストリ.md"
SHEETS_JSON = REPO_ROOT / "build" / "sheets_main.json"
DEFAULT_BOOKS = (
    REPO_ROOT / "dist" / "リスク提案ナビ_dev.xlsm",
    REPO_ROOT / "dist" / "リスク提案ナビ.xlsm",
)

SHEETS_KB_JSON = REPO_ROOT / "build" / "sheets_kb.json"
DEFAULT_KB_BOOKS = (REPO_ROOT / "dist" / "ナレッジブック.xlsx",)

# 15章§8.1「mock応答に現れるIDは…M-0012/L-03/S-0004/K-0003/P9/MC-0107のみ」。
# 17章T-14が「mock実行前にこれらの行がナレッジ雛形(T-03)へ投入済みであること」を
# 前提にしているため、KB照合(--kb)の一部としてここでも突合する(build_rpn.py
# --kb の自己検証と同じ台帳。二重実装ではなく、13章側とビルド側それぞれの
# 出口で同じ事実を検査する)。
KB_MOCK_ROW_IDS = {
    "メニュー一覧": "M-0012",
    "種目マスタ": "L-03",
    "型ライブラリ": "S-0004",
    "成功事例": "K-0003",
    "パターンマスタ": "P9",
    "機構ライブラリ": "MC-0107",
}

NAME_PREFIXES = ("hm_", "ci_", "hs_", "sp_", "gd_")


# ------------------------------------------------------------------------------
# Markdown の下ごしらえ
# ------------------------------------------------------------------------------
def _strip_marks(s: str) -> str:
    """強調(**)・コード引用(`)・全角空白を落とす。表セルの素の値を得るため。"""
    return s.replace("**", "").replace("`", "").replace("　", " ").strip()


def _drop_parens(s: str) -> str:
    """入れ子でない括弧の中身を落とす。13章の列セルは `detail(最大400字)` や
    `case_id(またはinbox_id)` のように補足を括弧で書くため、列名の抽出前に外す。
    括弧の中に "/" が入る例(E07xx受信箱/ウォッチ)があるので、分割より前に行う。"""
    prev = None
    while prev != s:
        prev = s
        s = re.sub(r"\([^()]*\)", "", s)
        s = re.sub(r"（[^（）]*）", "", s)
    return s


def _split_names(cell: str) -> list[str]:
    """「a / b / c」形式のセルを名前のリストへ割る。"""
    body = _drop_parens(_strip_marks(cell))
    return [p.strip() for p in body.split("/") if p.strip()]


def _parse_tables(lines: list[str]) -> list[dict]:
    """連続する `|` 行をMarkdown表として切り出す。区切り行(---)は読み飛ばす。
    戻り値: [{"header": [セル...], "rows": [[セル...], ...], "at": 行番号}]"""
    tables, i = [], 0
    while i < len(lines):
        if not lines[i].lstrip().startswith("|"):
            i += 1
            continue
        start = i
        block = []
        while i < len(lines) and lines[i].lstrip().startswith("|"):
            block.append(lines[i])
            i += 1
        if len(block) < 2:
            continue
        cells = [[c.strip() for c in row.strip().strip("|").split("|")] for row in block]
        header = cells[0]
        rows = [r for r in cells[2:]] if re.match(r"^[\s:|-]+$", block[1]) else cells[1:]
        tables.append({"header": header, "rows": rows, "at": start})
    return tables


# ------------------------------------------------------------------------------
# 13章のパース
# ------------------------------------------------------------------------------
class Ch13:
    """13章§2(本体ブック)から、シート・列・名前付きレンジ・configキーを取り出す。"""

    def __init__(self, text: str):
        self.text = text
        self.lines = text.split("\n")
        self.sheets: list[str] = []              # 13章が定義する本体シート名(節の出現順)
        self.flat_columns: dict[str, list[str]] = {}     # シート -> 列名(1テーブル)
        self.blocks: dict[str, list[tuple[str, list[str]]]] = {}  # シート -> [(ブロック, 列)]
        self.named_ranges: dict[str, list[str]] = {}     # シート -> 名前付きレンジ名
        self.config_keys: list[str] = []
        self.config_defaults: dict[str, str] = {}        # 機械比較できる行だけ
        self.config_unparsed: list[str] = []
        self.guard_sheet: str = ""
        self.form_sheets: list[str] = []
        self.table_sheets: list[str] = []
        self.prefix_of: dict[str, str] = {}      # シート -> 名前付きレンジ接頭辞
        self._parse()

    # -- 節の切り出し -----------------------------------------------------------
    def _sections(self):
        """`## 2.` 配下の `### 2.N ...` を (見出し, 行リスト) で返す。"""
        idx = [i for i, l in enumerate(self.lines) if l.startswith("## ")]
        start = next(i for i in idx if self.lines[i].startswith("## 2. "))
        end = next((i for i in idx if i > start), len(self.lines))
        body = self.lines[start:end]
        heads = [i for i, l in enumerate(body) if l.startswith("### ")]
        for k, h in enumerate(heads):
            tail = heads[k + 1] if k + 1 < len(heads) else len(body)
            yield body[h], body[h + 1:tail]

    def _parse(self):
        for head, body in self._sections():
            # 本体シートを定義する節は「### 2.N `シート名`...」の形。
            # 「### 2.4 ログ(...)」「### 2.8 企業ドシエファイル `<会社名>_...`」
            # 「### 2.9 画面シートの型と参照規約」はこの形にならないので外れる。
            m = re.match(r"^###\s+2\.\d+\s+`([^`]+)`", head)
            if m:
                self._parse_sheet_section(m.group(1), head, body)
                continue
            if "画面シートの型と参照規約" in head:
                self._parse_type_section(body)
            if re.match(r"^###\s+2\.\d+\s+ログ", head):
                self._parse_log_section(body)

    # -- 各節 -------------------------------------------------------------------
    def _parse_sheet_section(self, sheet: str, head: str, body: list[str]):
        self.sheets.append(sheet)
        tables = _parse_tables(body)

        if sheet == "config":
            self._parse_config(tables)
            return

        # 名前付きレンジ(帳票型と hs_ / sp_)。表・地の文をまたいで拾い、
        # 「ci_recipe_01 - ci_recipe_14」のような範囲表記は展開する。
        # 拾う接頭辞は13章§2.9の命名規約表から引く。節をまたぐ相互参照
        # (§2.16の本文が案件入力の ci_paste_hearing_answers_1 に触れる等)を
        # そのシートの持ち物と誤認しないため、自シートの接頭辞だけを拾う。
        prefix = self.prefix_of.get(sheet)
        if prefix:
            names = self._scan_named_ranges("\n".join(body), prefix)
            if names:
                self.named_ranges[sheet] = names

        # 列定義: 見出しの先頭セルが「列」の表だけを見る。直前に現れた
        # 「**ブロック `name`**」がそのブロック名になる(無ければ1シート1テーブル)。
        block_at = {}
        for i, line in enumerate(body):
            bm = re.search(r"\*\*ブロック\s+`([^`]+)`\*\*", line)
            if bm:
                block_at[i] = bm.group(1)

        found_blocks: list[tuple[str, list[str]]] = []
        for t in tables:
            if _strip_marks(t["header"][0]) != "列":
                continue
            cols: list[str] = []
            for row in t["rows"]:
                cols.extend(_split_names(row[0]))
            owners = [b for i, b in sorted(block_at.items()) if i < t["at"]]
            if owners:
                found_blocks.append((owners[-1], cols))
            else:
                if sheet in self.flat_columns:
                    self.flat_columns[sheet].extend(cols)
                else:
                    self.flat_columns[sheet] = cols
        if found_blocks:
            self.blocks[sheet] = found_blocks

    def _scan_named_ranges(self, text: str, prefix: str) -> list[str]:
        found: list[str] = []
        for m in re.finditer(r"\b([a-z]{2}_[a-z0-9_]+)\b", text):
            nm = m.group(1)
            if nm.startswith(prefix) and nm not in found:
                found.append(nm)
        # 「ci_recipe_01 - ci_recipe_14」形式の範囲を展開する。
        for m in re.finditer(r"\b([a-z][a-z0-9_]*_)(\d{2})\s*-\s*\1(\d{2})\b", text):
            stem, lo, hi = m.group(1), int(m.group(2)), int(m.group(3))
            if not stem.startswith(prefix):
                continue
            for n in range(lo, hi + 1):
                nm = f"{stem}{n:02d}"
                if nm not in found:
                    found.append(nm)
        return found

    def _parse_config(self, tables):
        for t in tables:
            if _strip_marks(t["header"][0]) != "name":
                continue
            for row in t["rows"]:
                keys = _split_names(row[0])
                self.config_keys.extend(keys)
                raw = _strip_marks(row[1]) if len(row) > 1 else ""
                toks = [p.strip() for p in raw.split("/")]
                if len(keys) == len(toks):
                    for k, v in zip(keys, toks):
                        self.config_defaults[k] = "" if v in ("(空)", "（空）") else v
                else:
                    # 既定値セルが自由記述(URL・パス・条件分岐など)で機械分割できない行。
                    # 黙って通さず、比較対象外であることを申告する。
                    self.config_unparsed.extend(keys)

    def _parse_log_section(self, body: list[str]):
        for line in body:
            m = re.match(r"^-\s+`(\w+)`(?:（[^）]*）)?\s*[:：]\s*(.+)$", line.strip())
            if not m:
                continue
            self.sheets.append(m.group(1))
            self.flat_columns[m.group(1)] = _split_names(m.group(2))

    def _parse_type_section(self, body: list[str]):
        blob = "\n".join(body)
        # 命名規約: 「`hm_`（HOME）/ `ci_`（案件入力）」「`hs_`（`ヒアリングシート`。§2.16）」
        for m in re.finditer(r"`([a-z]{2}_)`（`?([^）`。]+)`?", blob):
            self.prefix_of[_strip_marks(m.group(2))] = m.group(1)
        m = re.search(r"\*\*ガードシート\s+`([^`]+)`\*\*", blob)
        if m:
            self.guard_sheet = m.group(1)
            self.sheets.insert(0, self.guard_sheet)
        for t in _parse_tables(body):
            if _strip_marks(t["header"][0]) != "型":
                continue
            for row in t["rows"]:
                kind = _strip_marks(row[0])
                targets = re.findall(r"`([^`]+)`", row[1])
                if kind == "帳票型":
                    self.form_sheets = targets
                elif kind == "テーブル型":
                    self.table_sheets = targets


# ------------------------------------------------------------------------------
# 13章§3(ナレッジブック)のパース
# ------------------------------------------------------------------------------
# §3の列挙は§2と違い一枚岩でない: §3.2/§3.4/§3.10は§2と同じ「列|説明」の
# Markdown表だが、§3.3/§3.5/§3.6/§3.7/§3.8/§3.9は表を作らず1行の箇条書き
# (`a / b(注記) / c` 形式)で列を並べる。さらに§3.5-§3.8の4シートは、その
# 箇条書きの中に enum の値そのものを「/」区切りで割り込ませる書き方をする
# (例: `layer enum: detect（検知）/ prevent（予防）/ .../ underwrite（引受主体）
# / source_uid（...）`)。列名と enum 値をどちらも「/」区切りの同じ並びに書いて
# いるため、素朴に「/」で割ると enum 値が別の列であるかのように誤読する。
#
# 対処: 19章§3のenumレジストリが持つ「値の個数」を唯一の手がかりにする。
# 13章の箇条書きにだけ現れる4つの enum 語(layer/kind/status/classification)を
# 19章の対応する項目キー(mech.layer 等)へ写像し、そのレジストリ行の値の個数
# ぶんだけ後続トークンを列名として扱わず読み飛ばす。個数が特定できない語が
# 出てきたら黙って誤読せず例外にする(15章§0原則3と同じ「実在制約」の精神)。
_BULLET_ENUM_KEY = {
    "layer": "mech.layer",
    "kind": "watch.source_kind",
    "status": "target.status",
    "classification": "watch.classification",
}


def _parse_ch19_enum_lengths(text: str) -> dict[str, int]:
    """19章の全Markdown表から `項目キー -> enum値の個数` を拾う。項目キーは
    `mech.layer` のようなドット表記(19章§3の1列目)。値の個数が分かる行だけを
    拾うので、無関係な表の行は自然に無視される。"""
    out: dict[str, int] = {}
    for line in text.split("\n"):
        if not line.strip().startswith("|"):
            continue
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if len(cells) < 2:
            continue
        key = _drop_parens(_strip_marks(cells[0]))
        if not re.fullmatch(r"[a-zA-Z_][a-zA-Z0-9_.]*", key):
            continue
        vals = cells[1]
        if "/" not in vals:
            continue
        out[key] = len([v for v in vals.split("/") if v.strip()])
    return out


def _parse_bullet_columns(raw_line: str, enum_lengths: dict[str, int], sheet: str) -> list[str]:
    """箇条書き形式(表でない)1行の列挙を列名のリストへ割る。_BULLET_ENUM_KEY の
    語が出てきたら、19章から引いた個数ぶんの値トークンをまとめて読み飛ばす。"""
    body = _drop_parens(_strip_marks(raw_line))
    tokens = [t.strip() for t in body.split("/") if t.strip()]
    cols: list[str] = []
    i = 0
    while i < len(tokens):
        m = re.match(r"^(\w+)\s+enum:\s*\S", tokens[i])
        if m:
            colname = m.group(1)
            reg_key = _BULLET_ENUM_KEY.get(colname)
            n = enum_lengths.get(reg_key) if reg_key else None
            if not n:
                raise RuntimeError(
                    f"13章§3 '{sheet}': 箇条書き中の enum 語 '{colname}' の値数を"
                    f"19章§3レジストリ(項目キー={reg_key!r})から特定できません。"
                    "sheet_check.py の _BULLET_ENUM_KEY か19章の該当行を確認してください。")
            cols.append(colname)
            i += n
        else:
            cols.append(tokens[i])
            i += 1
    return cols


class Ch13Kb:
    """13章§3(ナレッジブック)から、シート・列を取り出す。"""

    def __init__(self, text: str, enum_lengths: dict[str, int]):
        self.lines = text.split("\n")
        self._enum_lengths = enum_lengths
        self.sheets: list[str] = []             # 出現順(§3.1の6シート＋§3.2-3.10)
        self.columns: dict[str, list[str]] = {}
        self.unparsed: list[str] = []            # 列を機械抽出できなかったシート
        self._parse()

    def _section3_lines(self) -> list[str]:
        idx = [i for i, l in enumerate(self.lines) if l.startswith("## ")]
        start = next(i for i in idx if re.match(r"^## 3\.\s", self.lines[i]))
        end = next((i for i in idx if i > start), len(self.lines))
        return self.lines[start:end]

    def _subsections(self):
        body = self._section3_lines()
        heads = [i for i, l in enumerate(body) if l.startswith("### ")]
        for k, h in enumerate(heads):
            tail = heads[k + 1] if k + 1 < len(heads) else len(body)
            yield body[h], body[h + 1:tail]

    def _parse(self):
        for head, body in self._subsections():
            if re.match(r"^###\s+3\.1\b", head):
                self._parse_31(body)
                continue
            m = re.match(r"^###\s+3\.\d+\s+`([^`]+)`", head)
            if not m:
                continue
            sheet = m.group(1)
            self.sheets.append(sheet)
            tables = _parse_tables(body)
            table = next((t for t in tables if _strip_marks(t["header"][0]) == "列"), None)
            if table is not None:
                cols: list[str] = []
                for row in table["rows"]:
                    cols.extend(_split_names(row[0]))
                self.columns[sheet] = cols
                continue
            # 表でない箇条書き形式(§3.3/3.5/3.6/3.7/3.8/3.9): 見出し直下の
            # 最初の非空行が列挙行(13章はこの1行の後に空行を置いてから次見出し
            # へ進む書式で統一されている)。
            line = next((l for l in body if l.strip()), None)
            if line is None:
                self.unparsed.append(sheet)
                continue
            self.columns[sheet] = _parse_bullet_columns(line, self._enum_lengths, sheet)

    def _parse_31(self, body: list[str]):
        """§3.1「既存6シート」: `- \\`name\\`: col / col / ...` の箇条書きが
        複数本並び、1行に2シート分(種目マスタ。メニュー種目対応)が「。」で
        連結される箇所がある。丸括弧を先に落としてから「。」で割るので、
        括弧内の説明文にある句点で誤って割れることはない。**`_strip_marks`は
        使わない**(バッククォートまで剥がすため、シート名を囲む `` ` `` が
        消えて下の正規表現が当たらなくなる。太字 `**` だけ手で剥がす)。"""
        text = _drop_parens("\n".join(body)).replace("**", "")
        for raw_line in text.split("\n"):
            line = raw_line.strip()
            if not line.startswith("-"):
                continue
            line = line[1:].strip()
            for chunk in line.split("。"):
                chunk = chunk.strip()
                m = re.match(r"^`([^`]+)`\s*[:：]\s*(.+)$", chunk)
                if not m:
                    continue
                sheet, rest = m.group(1), m.group(2)
                cols = [c.strip() for c in rest.split("/") if c.strip()]
                if not cols:
                    continue
                self.sheets.append(sheet)
                self.columns[sheet] = cols


# ------------------------------------------------------------------------------
# 台帳(build/sheets_kb.json)側の読み出し
# ------------------------------------------------------------------------------
class KbLedger:
    def __init__(self, data: dict):
        self.data = data
        self.specs = data["sheets"]
        self.order = [s["name"] for s in self.specs]
        self.by_name = {s["name"]: s for s in self.specs}

    def columns(self, name):
        s = self.by_name.get(name) or {}
        return [c["name"] for c in (s.get("columns") or [])] or None

    def seed_ids(self, name):
        """seed_rows の1列目(=各シートのID列)の値一覧。mock行の存在確認に使う。"""
        s = self.by_name.get(name) or {}
        return [row[0] for row in (s.get("seed_rows") or [])]


# ------------------------------------------------------------------------------
# ブック(dist/ナレッジブック.xlsx)側の読み出し
# ------------------------------------------------------------------------------
class KbBook:
    def __init__(self, path: Path):
        import openpyxl
        self.path = path
        self.wb = openpyxl.load_workbook(path, data_only=False)

    def header_row(self, sheet: str, width: int) -> list[str]:
        ws = self.wb[sheet]
        return [ws.cell(row=1, column=i).value for i in range(1, width + 1)]

    def column_a_from(self, sheet: str, start: int = 2) -> list[str]:
        ws = self.wb[sheet]
        out, r = [], start
        while ws.cell(row=r, column=1).value not in (None, ""):
            out.append(ws.cell(row=r, column=1).value)
            r += 1
        return out


# ------------------------------------------------------------------------------
# 台帳(sheets_main.json)側の読み出し
# ------------------------------------------------------------------------------
class Ledger:
    def __init__(self, data: dict):
        self.data = data
        self.specs = [s for s in data["sheets"] if not s.get("build_infrastructure")]
        self.infra = [s for s in data["sheets"] if s.get("build_infrastructure")]
        self.order = [s["name"] for s in self.specs]
        self.by_name = {s["name"]: s for s in self.specs}

    def flat_columns(self, name):
        s = self.by_name.get(name) or {}
        if s.get("role") == "config":
            return list(s.get("headers") or [])
        return [c["name"] for c in (s.get("columns") or [])] or None

    def blocks(self, name):
        s = self.by_name.get(name) or {}
        return [(b["name"], [c["name"] for c in b["columns"]])
                for b in (s.get("blocks") or [])] or None

    def named_ranges(self, name):
        s = self.by_name.get(name) or {}
        out = []
        for sec in s.get("sections") or []:
            out += [f["range"] for f in sec.get("fields") or []]
            if sec.get("anchor_table"):
                out.append(sec["anchor_table"]["range"])
        out += [f["range"] for f in s.get("header_fields") or []]
        return out

    def config_keys(self):
        s = self.by_name.get("config") or {}
        return [d["name"] for d in s.get("defaults") or []]

    def config_defaults(self):
        s = self.by_name.get("config") or {}
        return {d["name"]: d.get("value") for d in s.get("defaults") or []}

    def input_rule_omissions(self):
        """入力規則の省略台帳(17章T-46/裁定書4 項目17)を (sheet, target, enum) の
        集合と、宣言の整形不良の一覧で返す。"""
        raw = self.data.get("input_rule_omissions") or []
        decl, malformed = set(), []
        for i, e in enumerate(raw):
            if not all(k in e for k in ("sheet", "target", "enum")):
                malformed.append(f"エントリ{i}に sheet/target/enum のいずれかがありません: {e}")
                continue
            decl.add((e["sheet"], e["target"], e["enum"]))
        return decl, malformed

    def dv_fields(self):
        """台帳内の dv 指定のある列/フィールドを (sheet, target, enum, kind) で列挙する。
        target は flat/anchor=列物理名 / block='ブロック名.列名' / form・header=名前付きレンジ名。
        kind は 'flat'/'block'/'form' で、成果物側のセル解決の仕方を分ける。"""
        out = []
        for s in self.specs:
            sheet = s["name"]
            for i, c in enumerate(s.get("columns") or [], start=1):
                if c.get("dv"):
                    out.append((sheet, c["name"], c["dv"], ("flat", i)))
            for b in s.get("blocks") or []:
                for i, c in enumerate(b["columns"], start=1):
                    if c.get("dv"):
                        out.append((sheet, f"{b['name']}.{c['name']}", c["dv"],
                                    ("block", b["name"], i)))
            fields = []
            for sec in s.get("sections") or []:
                fields += sec.get("fields") or []
            fields += s.get("header_fields") or []
            for f in fields:
                if f.get("dv"):
                    out.append((sheet, f["range"], f["dv"], ("form", f["range"])))
        return out


# ------------------------------------------------------------------------------
# 比較の道具
# ------------------------------------------------------------------------------
class Report:
    def __init__(self):
        self.fails: list[str] = []
        self.warns: list[str] = []
        self.oks = 0

    def check(self, ok: bool, label: str, detail: str = ""):
        if ok:
            self.oks += 1
        else:
            self.fails.append(f"{label}: {detail}" if detail else label)
        return ok

    def warn(self, msg: str):
        self.warns.append(msg)

    def eq_seq(self, label, want, got):
        if want == got:
            self.oks += 1
            return True
        extra = [x for x in (got or []) if x not in (want or [])]
        missing = [x for x in (want or []) if x not in (got or [])]
        d = []
        if missing:
            d.append(f"不足={missing}")
        if extra:
            d.append(f"余分={extra}")
        if not d:
            d.append(f"順序違い 期待={want} 実際={got}")
        self.fails.append(f"{label}: " + " / ".join(d))
        return False

    def eq_set(self, label, want, got):
        return self.eq_seq(label, sorted(set(want)), sorted(set(got)))


def _norm_default(v) -> str:
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, float) and float(v).is_integer():
        return str(int(v))
    if isinstance(v, (int, float)):
        return str(v)
    s = "" if v is None else str(v).strip()
    return "" if s in ("(空)", "（空）") else s


# ------------------------------------------------------------------------------
# ブック側の読み出し
# ------------------------------------------------------------------------------
class Book:
    def __init__(self, path: Path):
        import openpyxl
        self.path = path
        self.wb = openpyxl.load_workbook(path, keep_vba=True, data_only=False)
        self.names = {}
        for nm, dn in dict(self.wb.defined_names).items():
            dests = list(dn.destinations)
            if len(dests) == 1:
                sheet, coord = dests[0]
                self.names[nm] = (sheet, coord.replace("$", ""))

    def header_row(self, sheet: str, row: int, width: int) -> list[str]:
        ws = self.wb[sheet]
        return [ws.cell(row=row, column=i).value for i in range(1, width + 1)]

    def row_of(self, coord: str) -> int:
        return int(re.sub(r"[^0-9]", "", coord))

    def column_a(self, sheet: str, start: int = 2) -> list[str]:
        ws = self.wb[sheet]
        out, r = [], start
        while ws.cell(row=r, column=1).value not in (None, ""):
            out.append(ws.cell(row=r, column=1).value)
            r += 1
        return out


# ------------------------------------------------------------------------------
# 入力規則の省略台帳(裁定書4 項目17)の実測ヘルパ
# ------------------------------------------------------------------------------
def _book_list_dv_cells(ws) -> set:
    """成果物ブックのワークシートで list 型の入力規則が覆う (col_idx, row) の集合。"""
    cells: set = set()
    try:
        dvs = ws.data_validations.dataValidation
    except Exception:
        return cells
    for dv in dvs:
        if getattr(dv, "type", None) != "list":
            continue
        for rng in dv.sqref.ranges:
            for c in range(rng.min_col, rng.max_col + 1):
                for r in range(rng.min_row, rng.max_row + 1):
                    cells.add((c, r))
    return cells


def _coord_col_row(coord: str):
    m = re.match(r"([A-Za-z]+)(\d+)", coord)
    if not m:
        return None
    col = 0
    for ch in m.group(1).upper():
        col = col * 26 + (ord(ch) - 64)
    return col, int(m.group(2))


def _measured_dv_omissions(led: "Ledger", book) -> set:
    """成果物ブックを実測し、台帳で dv 指定のある列/フィールドのうち、ブックに list
    入力規則が実際に付いていないもの(=ビルドが省略したもの)を (sheet, target, enum)
    の集合で返す。ビルドが入力規則を省略する唯一の理由は255字上限なので、この差分が
    そのまま『ビルド実測の省略一覧』になる。"""
    cache: dict = {}
    measured: set = set()
    for sheet, target, enum, kind in led.dv_fields():
        if sheet not in book.wb.sheetnames:
            continue
        ws = book.wb[sheet]
        covered = cache.get(sheet)
        if covered is None:
            covered = cache[sheet] = _book_list_dv_cells(ws)
        cell = None
        if kind[0] == "flat":
            cell = (kind[1], 2)                       # 列index, 先頭データ行=2
        elif kind[0] == "block":
            loc = book.names.get(kind[1])             # ブロックアンカーの名前付きレンジ
            if loc and loc[0] == sheet:
                cell = (kind[2], book.row_of(loc[1]) + 1)
        elif kind[0] == "form":
            loc = book.names.get(kind[1])             # 値セルの名前付きレンジ
            if loc and loc[0] == sheet:
                cr = _coord_col_row(loc[1])
                if cr:
                    cell = cr
        if cell is None:
            # ブック側でセルを特定できない=名前付きレンジ不足。省略判定へは巻き込まず、
            # [3][4]の名前付きレンジ検査側の失敗に委ねる。
            continue
        if cell not in covered:
            measured.add((sheet, target, enum))
    return measured


# ------------------------------------------------------------------------------
# 本体
# ------------------------------------------------------------------------------
def run(book_path: Path | None, rep: Report) -> None:
    ch13 = Ch13(SPEC_CH13.read_text(encoding="utf-8"))
    led = Ledger(json.loads(SHEETS_JSON.read_text(encoding="utf-8")))
    book = Book(book_path) if book_path else None

    print(f"13章:   {SPEC_CH13}")
    print(f"台帳:   {SHEETS_JSON}")
    print(f"ブック: {book_path if book_path else '(--no-book のため未検査)'}")
    print()

    # --- 1. シート集合 --------------------------------------------------------
    print("[1] シート集合(13章§2の節見出し＋§2.4のログ＋§2.9のガードシート)")
    print(f"    13章が定義する本体シート: {len(ch13.sheets)}枚")
    if led.infra:
        print("    仕様外のビルド機構シート(13章に存在しないため照合対象から明示的に除外): "
              + ", ".join(s["name"] for s in led.infra))
    rep.eq_set("13章 ⇔ 台帳のシート集合", ch13.sheets, led.order)
    if book:
        got = [n for n in book.wb.sheetnames
               if n not in {s["name"] for s in led.infra}]
        rep.eq_set("13章 ⇔ ブックのシート集合", ch13.sheets, got)
        rep.eq_seq("台帳 ⇔ ブックのシート順序", led.order, got)
    print(f"    ガードシート: {ch13.guard_sheet or '(13章§2.9から取得できず)'}")
    rep.check(bool(ch13.guard_sheet), "13章§2.9のガードシート規定を読めること")
    if ch13.guard_sheet:
        gs = led.by_name.get(ch13.guard_sheet)
        rep.check(gs is not None and gs.get("role") == "guard",
                  "ガードシートの role", f"台帳 role={gs and gs.get('role')}")
        rep.check(gs is not None and gs.get("state") == "visible",
                  "ガードシートの初期可視状態", "13章§2.9は配布ビルドで可視・先頭")
        if book:
            rep.check(book.wb.sheetnames[0] == ch13.guard_sheet,
                      "ガードシートが先頭", f"実際={book.wb.sheetnames[0]}")
            rep.check(book.wb[ch13.guard_sheet].sheet_state == "visible",
                      "ガードシートがブックで可視",
                      f"実際={book.wb[ch13.guard_sheet].sheet_state}")

    # --- 2. シート型 ----------------------------------------------------------
    print("\n[2] シート型(13章§2.9の型表 ⇔ 台帳の role)")
    print(f"    帳票型={ch13.form_sheets} / テーブル型={ch13.table_sheets}")
    for nm in ch13.form_sheets:
        s = led.by_name.get(nm) or {}
        rep.check(s.get("role") == "form", f"'{nm}' は帳票型", f"台帳 role={s.get('role')}")
        rep.check(not s.get("columns") and not s.get("blocks"),
                  f"'{nm}' は1行目ヘッダを持たない", "帳票型に columns/blocks を置かない")
    for nm in ch13.table_sheets:
        s = led.by_name.get(nm) or {}
        rep.check(s.get("role") == "table", f"'{nm}' はテーブル型", f"台帳 role={s.get('role')}")

    # --- 3. 列物理名 ----------------------------------------------------------
    print("\n[3] 列物理名と物理順(テーブル型)")
    for sheet in ch13.sheets:
        want_flat = ch13.flat_columns.get(sheet)
        want_blocks = ch13.blocks.get(sheet)
        if want_flat:
            got = led.flat_columns(sheet)
            rep.eq_seq(f"[台帳] {sheet} の列", want_flat, got)
            print(f"    {sheet}: {len(want_flat)}列")
            if book and sheet in book.wb.sheetnames:
                bgot = book.header_row(sheet, 1, len(want_flat))
                rep.eq_seq(f"[ブック] {sheet} の1行目ヘッダ", want_flat, bgot)
        if want_blocks:
            got = led.blocks(sheet) or []
            rep.eq_seq(f"[台帳] {sheet} のブロック順",
                       [b for b, _ in want_blocks], [b for b, _ in got])
            gmap = dict(got)
            print(f"    {sheet}: {len(want_blocks)}ブロック / "
                  f"{sum(len(c) for _, c in want_blocks)}列")
            for bname, cols in want_blocks:
                rep.eq_seq(f"[台帳] {sheet}/{bname} の列", cols, gmap.get(bname))
                if book and sheet in book.wb.sheetnames:
                    loc = book.names.get(bname)
                    if not rep.check(loc is not None,
                                     f"[ブック] ブロックアンカー '{bname}' の名前付きレンジ",
                                     "13章§2.9: ブロック名と同名の名前付きレンジで"
                                     "ヘッダ行の左端セルを指す"):
                        continue
                    if not rep.check(loc[0] == sheet,
                                     f"[ブック] '{bname}' の所在シート",
                                     f"実際={loc[0]}"):
                        continue
                    bgot = book.header_row(sheet, book.row_of(loc[1]), len(cols))
                    rep.eq_seq(f"[ブック] {sheet}/{bname} のヘッダ行", cols, bgot)

    # --- 4. 名前付きレンジ ----------------------------------------------------
    print("\n[4] 名前付きレンジ(帳票型 hm_/ci_/gd_ と 見出し用 hs_/sp_)")
    total_hs_sp = 0
    for sheet, want in ch13.named_ranges.items():
        got = led.named_ranges(sheet)
        rep.eq_set(f"[台帳] {sheet} の名前付きレンジ", want, got)
        n_hs_sp = sum(1 for n in want if n.startswith(("hs_", "sp_")))
        total_hs_sp += n_hs_sp
        print(f"    {sheet}: {len(want)}本" + (f"(うち hs_/sp_ {n_hs_sp}本)" if n_hs_sp else ""))
        if book:
            bgot = [n for n, (sh, _) in book.names.items() if sh == sheet
                    and n.startswith(NAME_PREFIXES)]
            rep.eq_set(f"[ブック] {sheet} の名前付きレンジ", want, bgot)
    rep.check(total_hs_sp == 8,
              "hs_/sp_ の見出し用名前付きレンジは計8本(19章§5)",
              f"13章から読めたのは{total_hs_sp}本")

    # --- 5. configキー --------------------------------------------------------
    print("\n[5] configキー(13章§2.3)")
    print(f"    13章のキー数: {len(ch13.config_keys)}")
    rep.eq_seq("[台帳] configキー(順序含む)", ch13.config_keys, led.config_keys())
    if book and "config" in book.wb.sheetnames:
        rep.eq_seq("[ブック] configキー(順序含む)", ch13.config_keys, book.column_a("config"))
    lv = led.config_defaults()
    cmp_n = 0
    for k, want in ch13.config_defaults.items():
        if k not in lv:
            continue
        cmp_n += 1
        rep.check(_norm_default(lv[k]) == _norm_default(want),
                  f"[台帳] config既定値 {k}",
                  f"13章={want!r} 台帳={_norm_default(lv[k])!r}")
    print(f"    既定値を機械比較したキー: {cmp_n} / 比較対象外(13章の既定値欄が自由記述): "
          f"{sorted(set(ch13.config_unparsed))}")
    if book and "config" in book.wb.sheetnames:
        ws = book.wb["config"]
        for r in range(2, 2 + len(ch13.config_keys)):
            k = ws.cell(row=r, column=1).value
            if k not in lv:
                continue
            got = ws.cell(row=r, column=2).value
            if k == "mock_llm":
                # --dev / --prod で値が変わる唯一のキー。型だけを見る。
                rep.check(isinstance(got, bool), "[ブック] mock_llm は真偽値",
                          f"実際={got!r}")
                continue
            rep.check(_norm_default(got) == _norm_default(lv[k]),
                      f"[ブック] config既定値 {k}",
                      f"台帳={_norm_default(lv[k])!r} ブック={_norm_default(got)!r}")

    # --- 6. 19章§4との相互確認(参考。19章§5チェックリストの追随漏れ検知) ------
    print("\n[6] 19章§4との相互確認(参考・WARNのみ。19章§5は13章と19章の同時更新を求める)")
    _cross_check_ch19(ch13, rep)

    # --- 7. 入力規則の省略台帳(17章T-46・裁定書4 項目17) -----------------------
    print("\n[7] 入力規則の省略台帳(255字上限で省略した列の宣言 ⇔ ビルド実測の一致)")
    decl, malformed = led.input_rule_omissions()
    for msg in malformed:
        rep.check(False, "input_rule_omissions の宣言整形", msg)
    enums = led.data.get("enums") or {}
    for sheet, target, enum in sorted(decl):
        rep.check(sheet in led.by_name, f"省略台帳: シート '{sheet}' が台帳に実在",
                  f"target={target}")
        rep.check(enum in enums, f"省略台帳: enum '{enum}' が enums に実在",
                  f"{sheet}!{target}")
    print(f"    宣言された省略: {len(decl)}件 {sorted(decl)}")
    if book:
        measured = _measured_dv_omissions(led, book)
        print(f"    ビルド実測の省略(dv指定列で入力規則が付いていないもの): "
              f"{len(measured)}件 {sorted(measured)}")
        rep.eq_set("入力規則の省略台帳 ⇔ ビルド実測(成果物ブック)", decl, measured)
    else:
        rep.warn("--no-book のため入力規則の省略台帳とビルド実測の一致検査を省略しました")


def _cross_check_ch19(ch13: Ch13, rep: Report) -> None:
    if not SPEC_CH19.exists():
        rep.warn("19章が見つからないため相互確認を省略しました")
        return
    text = SPEC_CH19.read_text(encoding="utf-8")
    m = re.search(r"\*\*本体シート\*\*（\*\*全(\d+)枚\*\*）[:：]\s*(.+)", text)
    if m:
        n = int(m.group(1))
        # 一覧の後ろに地の文(「。シート型（…）と列・…の定義は13章§2.9-§2.17」)が
        # 続くため、括弧を落としてから最初の句点で切る。
        listed = [_strip_marks(x) for x in
                  re.split(r"\s*/\s*", _drop_parens(m.group(2)).split("。")[0])]
        listed = [x for x in listed if x]
        if n != len(ch13.sheets):
            rep.warn(f"19章§4の本体シート枚数({n})と13章§2から導いた枚数"
                     f"({len(ch13.sheets)})が違います")
        diff = sorted(set(listed) ^ set(ch13.sheets))
        if diff:
            rep.warn(f"19章§4の本体シート一覧と13章§2の差分: {diff}")
    else:
        rep.warn("19章§4の本体シート一覧を読み取れませんでした")

    m = re.search(r"\*\*configキー\*\*[^（(]*[（(]([^）)]*)[）)]", text)
    if m:
        listed = [x.strip() for x in m.group(1).split("/") if x.strip()]
        listed = [x for x in listed if re.fullmatch(r"[a-z0-9_]+", x)]
        diff = sorted(set(listed) ^ set(ch13.config_keys))
        if diff:
            rep.warn(f"19章§4のconfigキー一覧と13章§2.3の差分: {diff}")
    else:
        rep.warn("19章§4のconfigキー一覧を読み取れませんでした")


# ------------------------------------------------------------------------------
# 本体(--kb: ナレッジブック)
# ------------------------------------------------------------------------------
def run_kb(book_path: Path | None, rep: Report) -> None:
    ch19_text = SPEC_CH19.read_text(encoding="utf-8") if SPEC_CH19.exists() else ""
    enum_lengths = _parse_ch19_enum_lengths(ch19_text)
    ch13kb = Ch13Kb(SPEC_CH13.read_text(encoding="utf-8"), enum_lengths)
    led = KbLedger(json.loads(SHEETS_KB_JSON.read_text(encoding="utf-8")))
    book = KbBook(book_path) if book_path else None

    print(f"13章:   {SPEC_CH13} (§3 ナレッジブック)")
    print(f"台帳:   {SHEETS_KB_JSON}")
    print(f"ブック: {book_path if book_path else '(--no-book のため未検査)'}")
    print()

    # --- 1. シート集合と順序 ----------------------------------------------
    print("[1] シート集合と順序(13章§3の節見出し＋§3.1の6シート)")
    print(f"    13章§3が定義するシート: {len(ch13kb.sheets)}枚")
    if ch13kb.unparsed:
        rep.warn(f"13章§3で列を機械抽出できなかったシート: {ch13kb.unparsed}"
                  "(パーサ要更新の疑い)")
    rep.eq_seq("13章§3 ⇔ 台帳のシート集合と順序", ch13kb.sheets, led.order)
    if book:
        got = list(book.wb.sheetnames)
        rep.eq_seq("13章§3 ⇔ ブックのシート集合と順序", ch13kb.sheets, got)

    # --- 2. 列物理名と物理順(全シート1シート1テーブル) ---------------------
    print("\n[2] 列物理名と物理順(全シート1シート1テーブル。13章§3にブロック縦積みは無い)")
    for sheet in ch13kb.sheets:
        want = ch13kb.columns.get(sheet)
        if not want:
            continue
        got = led.columns(sheet)
        rep.eq_seq(f"[台帳] {sheet} の列", want, got)
        print(f"    {sheet}: {len(want)}列")
        if book and sheet in book.wb.sheetnames:
            bgot = book.header_row(sheet, len(want))
            rep.eq_seq(f"[ブック] {sheet} の1行目ヘッダ", want, bgot)

    # --- 3. 業種マスタ30行(17章T-03のDoD) ----------------------------------
    print("\n[3] 業種マスタ30行(17章T-03のDoD)")
    seed_n = len(led.seed_ids("業種マスタ"))
    rep.check(seed_n >= 30, "[台帳] 業種マスタのseed_rowsが30行以上", f"実際={seed_n}行")
    print(f"    [台帳] 業種マスタ: {seed_n}行")
    if book and "業種マスタ" in book.wb.sheetnames:
        n = len(book.column_a_from("業種マスタ"))
        rep.check(n >= 30, "[ブック] 業種マスタの行数が30行以上", f"実際={n}行")
        print(f"    [ブック] 業種マスタ: {n}行")
    elif book:
        rep.check(False, "[ブック] 業種マスタシートが存在しません")

    # --- 4. mock用ナレッジ行6件(15章§8.1・17章T-14前提) --------------------
    print("\n[4] mock用ナレッジ行6件(15章§8.1・17章T-14前提)")
    for sheet, mid in KB_MOCK_ROW_IDS.items():
        rep.check(mid in led.seed_ids(sheet), f"[台帳] {sheet} に mock行 {mid}")
        if book and sheet in book.wb.sheetnames:
            rep.check(mid in [str(v) for v in book.column_a_from(sheet)],
                      f"[ブック] {sheet} に mock行 {mid}")
        elif book:
            rep.check(False, f"[ブック] シート'{sheet}'が存在しません({mid})")

    # --- 5. 19章§4との相互確認(参考。WARNのみ) -----------------------------
    print("\n[5] 19章§4との相互確認(参考・WARNのみ。19章§5は13章と19章の同時更新を求める)")
    _cross_check_ch19_kb(ch13kb, ch19_text, rep)


def _cross_check_ch19_kb(ch13kb: Ch13Kb, text: str, rep: Report) -> None:
    if not text:
        rep.warn("19章が見つからないため相互確認を省略しました")
        return
    m = re.search(r"\*\*ナレッジブックシート\*\*[:：]\s*(.+)", text)
    if not m:
        rep.warn("19章§4のナレッジブックシート一覧を読み取れませんでした")
        return
    listed = [_strip_marks(x) for x in
              re.split(r"\s*/\s*", _drop_parens(m.group(1)).split("。")[0])]
    listed = [x for x in listed if x]
    diff = sorted(set(listed) ^ set(ch13kb.sheets))
    if diff:
        rep.warn(f"19章§4のナレッジブックシート一覧と13章§3の差分: {diff}")


def main() -> int:
    ap = argparse.ArgumentParser(description="13章とシート台帳・ビルド済みブックの照合")
    ap.add_argument("--book", default=None,
                    help="検査するブック(既定: --kb 無指定なら dist/ の dev -> prod、"
                         "--kb 指定なら dist/ナレッジブック.xlsx)")
    ap.add_argument("--no-book", action="store_true", help="ブックを検査せず台帳だけ突合する")
    ap.add_argument("--kb", action="store_true",
                    help="本体ブックでなくナレッジブック(13章§3・17章T-03)を照合する")
    args = ap.parse_args()

    default_books = DEFAULT_KB_BOOKS if args.kb else DEFAULT_BOOKS
    book_path = None
    if not args.no_book:
        if args.book:
            book_path = Path(args.book)
            if not book_path.is_absolute():
                book_path = REPO_ROOT / book_path
            if not book_path.exists():
                print(f"ERROR: ブックが見つかりません: {book_path}", file=sys.stderr)
                return 1
        else:
            book_path = next((p for p in default_books if p.exists()), None)
            if book_path is None:
                build_cmd = "python3 build/build_rpn.py --kb" if args.kb \
                    else "python3 build/build_rpn.py --dev"
                print(f"ERROR: dist/ にビルド済みブックがありません。"
                      f"先に {build_cmd} を実行してください"
                      "(台帳だけを見るなら --no-book)。", file=sys.stderr)
                return 1

    rep = Report()
    if args.kb:
        print("=== sheet_check.py --kb (13章§3 ⇔ build/sheets_kb.json ⇔ ナレッジブック) ===")
        run_kb(book_path, rep)
    else:
        print("=== sheet_check.py (13章 ⇔ build/sheets_main.json ⇔ 成果物ブック) ===")
        run(book_path, rep)

    print("\n" + "=" * 78)
    if rep.warns:
        print(f"WARN {len(rep.warns)}件(章間の追随漏れの疑い。17章§1のとおり本タスクでは"
              f"docs/を直さず報告する):")
        for w in rep.warns:
            print(f"  - {w}")
    if rep.fails:
        print(f"FAIL {len(rep.fails)}件 / PASS {rep.oks}件")
        for f in rep.fails:
            print(f"  - {f}")
        src = "build/sheets_kb.json" if args.kb else "build/sheets_main.json"
        print(f"\n13章が正・台帳が従。ずれは {src} 側を直すこと"
              "(13章自体の誤りを疑ったときは直さずに報告する)。")
        return 1
    if args.kb:
        print(f"OK: 全{rep.oks}項目一致 "
              "(シート集合と順序・列名と物理順・業種マスタ30行・mock用ナレッジ行6件)")
    else:
        print(f"OK: 全{rep.oks}項目一致 "
              "(シート集合・シート型・列名と物理順・名前付きレンジ・configキーと既定値)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
